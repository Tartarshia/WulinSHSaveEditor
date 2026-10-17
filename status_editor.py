"""Player status catalog and guarded, atomic in-memory status editing."""
from copy import deepcopy
from io import BytesIO
import math
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


def tokens(value):
    if isinstance(value, list): return value
    return [x.strip() for x in str(value or "").split(",") if x.strip()]


def flag(value):
    return str(value).strip().casefold() in {"true", "1"}


def numeric_effects(row):
    result = {}
    for effect in tokens(row.get("continousEffect")):
        key, separator, value = str(effect).partition("|")
        if not separator or not key or "|" in value: raise ValueError("包含未支持的多参数效果")
        value = float(value)
        if not math.isfinite(value): raise ValueError("效果值不是有限数值")
        result[key] = result.get(key, 0) + value
    return result


def readonly_reason(row):
    if row.get("unknown"): return "当前游戏表中未找到定义，仅显示存档记录"
    if row.get("displayType") in {"经脉", "家园"}: return "由经脉或家园系统管理，只读"
    if str(row.get("uName", "")).startswith("剧情"): return "剧情联动状态，只读"
    key = str(row.get("uName", ""))
    if key.startswith(("团队体力效果_", "团队心情效果_")): return "由队伍体力或心情自动管理，只读"
    if key.startswith("主角与") and "情缘buff" in key: return "由情缘关系系统管理，只读"
    for field, reason in (("instantEffect", "有获得时的即时属性或事件效果"),
                          ("endRemoveEffect", "有移除时的属性或事件效果"),
                          ("dynamicModifier", "由动态脚本计算"), ("battleAruas", "包含战斗光环联动"),
                          ("tagsForActive", "需要其他标签才能生效"), ("tagsNotActive", "有互斥或免疫条件"),
                          ("tagKillOnActive", "激活会移除其他状态"), ("teammateBuff", "涉及队友联动")):
        if row.get(field): return reason + "，只读"
    if str(row.get("tickType", "")).casefold() not in {"infinite", "time", "battle", "battleaction"}:
        return "未支持的状态计时方式，只读"
    try:
        if int(row.get("lifeTime", 0)) <= 0: return "即时效果或非持续状态，只读"
        if any("@" in key for key in numeric_effects(row)): return "效果作用于团队或其他目标，只读"
    except (ValueError, TypeError): return "效果或时长格式无法安全解析，只读"
    return ""


def load_status_catalog(game_root: Path):
    from game_data import read_localized_item_info, plain_game_text
    import UnityPy
    tables = {}
    with ZipFile(game_root / "Wulin_Data/StreamingAssets/ModEditor/ModEditor.zip") as archive:
        for info in archive.infolist():
            if not info.filename.endswith(".xlsx") or "/xls/1/" not in info.filename.replace("\\", "/"): continue
            workbook = load_workbook(BytesIO(archive.read(info)), read_only=True, data_only=True)
            for name in ("BuffData", "TraitData"):
                if name not in workbook.sheetnames: continue
                sheet = workbook[name]
                headers = next(sheet.iter_rows(min_row=3, max_row=3, max_col=25, values_only=True))
                tables.setdefault(name, []).extend({key: value for key, value in zip(headers, row) if key}
                                                   for row in sheet.iter_rows(min_row=4, max_col=25, values_only=True) if row[2])
            workbook.close()
    if "BuffData" not in tables or "TraitData" not in tables: raise ValueError("游戏表中缺少状态或角色创建特征资料")
    names, descriptions = {}, {}
    env = UnityPy.load(str(game_root / "Wulin_Data/StreamingAssets/aa/StandaloneWindows64/config_assets_all.bundle"))
    for obj in env.objects:
        if obj.type.name == "MonoBehaviour" and obj.peek_name() == "LocData":
            raw = obj.get_raw_data()
            names = read_localized_item_info(raw, "BuffName_")
            descriptions = read_localized_item_info(raw, "BuffInfo_")
            break
    creation = {}
    for row in tables["TraitData"]:
        if flag(row.get("isUseByPlayer")):
            for key in tokens(row.get("bonusBuff")):
                creation.setdefault(key, []).append(str(row["uName"]))
    result = {}
    for index, row in enumerate(tables["BuffData"]):
        key = str(row["uName"])
        row = dict(row)
        row["isVisible"] = flag(row.get("isVisible"))
        row["positive"] = flag(row.get("isPositive"))
        for field in ("continousEffect", "instantEffect", "endRemoveEffect", "dynamicModifier", "battleAruas", "tagsForActive", "tagsNotActive", "tagKillOnActive", "teammateBuff"):
            row[field] = tokens(row.get(field))
        loc = str(row.get("locKey") or key)
        try: uid, rarity = int(row.get("uid", 0)), int(row.get("rarity", -1))
        except (ValueError, TypeError): uid, rarity = 0, -1
        row.update(id=uid, key=key, rarity=rarity, order=index,
                   name=plain_game_text(names.get(loc, key)),
                   description=plain_game_text(descriptions.get(loc, "游戏资源中暂无文字说明。")),
                   creation=creation.get(key, []),
                   source="创建角色可选" if key in creation else str(row.get("displayType") or "其他状态"),
                   permanent=str(row.get("tickType", "")).casefold() == "infinite")
        row["readonly"] = readonly_reason(row)
        if not row["isVisible"] and not row["creation"]:
            row["readonly"] = "游戏内部隐藏状态，只读"
        result[key] = row
    return result


def instance_reason(row, library, instances):
    reason = row.get("readonly", "")
    if reason: return reason
    if row["key"] in library.get("buffsFromTemplete", []): return "角色模板自带状态，只读"
    if any(x.get("overrideValues") is not None for x in instances): return "包含实例专属或动态数值，只读"
    try:
        effects = numeric_effects(row)
        for x in instances:
            stack = float(x.get("stack", 1))
            if not math.isfinite(stack) or stack < 0: raise ValueError()
        for key in effects:
            value = float(library.get("buffStats", {}).get(key, 0))
            if not math.isfinite(value): raise ValueError()
            if any(x.get("isValid") for x in instances) and key not in library.get("buffStats", {}):
                return "缺少已生效属性缓存，无法安全移除"
    except (ValueError, TypeError): return "状态实例或属性缓存格式异常，只读"
    return ""


def list_statuses(hero, catalog, query="", kind="all", source="all", sort="default"):
    library = hero.get("buffLib") or {}
    current = {}
    for instance in library.get("buffInstances", []):
        current.setdefault(instance.get("dataId"), []).append(instance)
    rows = []
    for key in dict.fromkeys([*catalog, *current]):
        detail = catalog.get(key) or dict(key=key, id=0, rarity=-1, order=len(catalog), name=str(key), source="未知状态",
                                         description="当前游戏导出表未包含此状态；保留原始记录，只读。", creation=[], permanent=None, readonly="当前游戏表中未找到定义，仅显示存档记录")
        if not detail.get("isVisible") and key not in current and not detail.get("creation"): continue
        instances = current.get(key, [])
        rows.append(dict(key=key, id=detail["id"], name=detail["name"], description=detail["description"],
                         source=detail["source"], creation=detail["creation"], rarity=detail["rarity"], order=detail["order"],
                         permanent=detail["permanent"], owned=bool(instances), active=any(x.get("isValid") for x in instances),
                         records=len(instances), remaining=[x.get("lifeTime") for x in instances],
                         lifetime=detail.get("lifeTime"), tick=detail.get("tickType", ""),
                         readonly=instance_reason(detail, library, instances)))
    sources = list(dict.fromkeys(["创建角色可选", *[x["source"] for x in rows]]))
    total, owned = len(rows), sum(x["owned"] for x in rows)
    needle = query.strip().casefold()
    rows = [x for x in rows if (not needle or needle in " ".join((x["name"], str(x["id"]), x["key"])).casefold())
            and (source == "all" or x["source"] == source)
            and (kind == "all" or kind == "permanent" and x["permanent"] is True or kind == "temporary" and x["permanent"] is False)]
    def order(row):
        if sort == "id": return row["id"], row["order"]
        if sort == "rarity": return -row["rarity"], row["order"]
        return (row["order"],)
    rows.sort(key=lambda row: (not row["owned"], *order(row)))
    return dict(rows=rows, total=total, owned=owned, sources=sources)


def set_status(hero, catalog, key, owned, current_quarter):
    if not isinstance(owned, bool): raise ValueError("获得状态必须是复选框布尔值")
    if key not in catalog: raise ValueError("未知状态不能修改")
    library = hero.get("buffLib")
    if not isinstance(library, dict) or not isinstance(library.get("buffInstances"), list) or not isinstance(library.get("buffStats"), dict):
        raise ValueError("存档状态库结构不完整")
    detail = catalog[key]
    instances = [x for x in library["buffInstances"] if x.get("dataId") == key]
    reason = instance_reason(detail, library, instances)
    if reason: raise ValueError(reason)
    if bool(instances) == owned: return
    now = float(current_quarter)
    if not math.isfinite(now) or now < 0: raise ValueError("存档游戏时间无效")
    replacement = deepcopy(library)
    effects = numeric_effects(detail)
    multiplier = 1 if owned else -sum(float(x.get("stack", 1)) for x in instances if x.get("isValid"))
    for name, value in effects.items():
        total = float(replacement["buffStats"].get(name, 0)) + value * multiplier
        if not math.isfinite(total): raise ValueError("状态属性计算超出范围")
        replacement["buffStats"][name] = total
    if owned:
        replacement["buffInstances"].append(dict(isValid=True, stack=1, initTime=now, lifeTime=int(detail["lifeTime"]),
                                                dataId=key, overrideValues=None, IsNew=True))
    else:
        replacement["buffInstances"] = [x for x in replacement["buffInstances"] if x.get("dataId") != key]
    hero["buffLib"] = replacement
